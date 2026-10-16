"""Fixed-layout DOCX export with native editable report text over page artwork.

A PDF's original paragraphs/tables cannot reliably be reconstructed. Render
artwork/photos/annotations as page backgrounds and keep entered values as
positioned Word text boxes. Reuse the PDF renderer for fit and coordinates.
"""

from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from lxml import etree
import pdfplumber
import pypdfium2 as pdfium

from .template_pdf import generate_pdf, generate_field_layers

V = "urn:schemas-microsoft-com:vml"
O = "urn:schemas-microsoft-com:office:office"


def _page_picture(paragraph, payload, width, height, identity):
    inline = (
        paragraph.add_run()
        .add_picture(BytesIO(payload), width=Pt(width), height=Pt(height))
        ._inline
    )
    # Turn an inline picture into an absolute, behind-text page anchor.
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "1",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis in ("H", "V"):
        position = OxmlElement("wp:position" + axis)
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    anchor.append(OxmlElement("wp:wrapNone"))
    properties = inline.find(qn("wp:docPr"))
    properties.set("id", str(identity))
    properties.set("name", f"Page {identity} background")
    properties.set(
        "descr", "Fixed template artwork, photographs and report annotations"
    )
    anchor.append(properties)
    anchor.append(inline.find(qn("wp:cNvGraphicFramePr")))
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)


def _text_box(paragraph, line, x, y, width, identity, field):
    chars = line["chars"]
    size = chars[0]["size"]
    font = {
        "Times-Roman": "Times New Roman",
        "Times-Bold": "Times New Roman",
        "Helvetica": "Arial",
        "Helvetica-Bold": "Arial",
        "Courier": "Courier New",
    }[field.font]
    pict = OxmlElement("w:pict")
    shape = etree.SubElement(pict, f"{{{V}}}rect", nsmap={"v": V, "o": O})
    shape.set("id", f"ReportValue_{identity}")
    shape.set("filled", "f")
    shape.set("stroked", "f")
    shape.set(
        "style",
        f"position:absolute;margin-left:{x:.3f}pt;margin-top:{y:.3f}pt;width:{width:.3f}pt;height:{size * 1.6:.3f}pt;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page",
    )
    textbox = etree.SubElement(shape, f"{{{V}}}textbox")
    textbox.set("inset", "0,0,0,0")
    content = OxmlElement("w:txbxContent")
    p = OxmlElement("w:p")
    properties = parse_xml(
        f'<w:pPr {nsdecls("w")}><w:spacing w:before="0" w:after="0" w:line="{round(size * 20 * 1.15)}" w:lineRule="exact"/><w:jc w:val="left"/></w:pPr>'
    )
    p.append(properties)
    run = OxmlElement("w:r")
    formatting = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), font)
    formatting.append(fonts)
    for tag in ("sz", "szCs"):
        element = OxmlElement("w:" + tag)
        element.set(qn("w:val"), str(round(size * 2)))
        formatting.append(element)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), field.color.lstrip("#"))
    formatting.append(color)
    if "Bold" in field.font:
        formatting.append(OxmlElement("w:b"))
    run.append(formatting)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = line["text"]
    run.append(text)
    p.append(run)
    content.append(p)
    textbox.append(content)
    paragraph.add_run()._r.append(pict)


def generate_docx(blank_pdf, definition, report, images):
    # This also applies the same validation/auto-fit rules as PDF downloads.
    layers = generate_field_layers(definition, report, images)
    artwork_definition = definition.model_copy(deep=True)
    artwork_report = report.model_copy(deep=True)
    for field in artwork_definition.fields:
        if field.kind not in {"image", "checkbox"}:
            field.default_value = ""
            if field.kind == "calculated":
                field.kind, field.calculation = "text", None
            artwork_report.values[field.id] = ""
        elif field.kind == "checkbox":
            artwork_report.values[field.id] = (
                report.values.get(field.id, field.default_value) is True
            )
    artwork = generate_pdf(blank_pdf, artwork_definition, artwork_report, images)
    document = Document()
    document.core_properties.title = report.title
    document.core_properties.author = ""
    document.core_properties.comments = "Fixed-layout export. Report values are editable Word text boxes; template artwork is not reconstructed as native Word tables."
    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(1)
    normal.font.size = Pt(1)
    page_paragraphs = []
    with pdfium.PdfDocument(artwork) as rendered:
        for index, page in enumerate(definition.pages):
            section = (
                document.sections[0]
                if index == 0
                else document.add_section(WD_SECTION_START.NEW_PAGE)
            )
            section.page_width, section.page_height = Pt(page.width), Pt(page.height)
            section.left_margin = section.right_margin = section.top_margin = (
                section.bottom_margin
            ) = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
            paragraph = document.add_paragraph()
            page_paragraphs.append(paragraph)
            picture = BytesIO()
            rendered[index].render(
                scale=min(2, 2400 / max(page.width, page.height))
            ).to_pil().save(picture, format="PNG")
            _page_picture(
                paragraph, picture.getvalue(), page.width, page.height, index + 1
            )
    if layers["pdf"]:
        import base64

        fields = {field.id: field for field in definition.fields}
        with pdfplumber.open(BytesIO(base64.b64decode(layers["pdf"]))) as values_pdf:
            identity = 0
            for layer_page, field_id in zip(values_pdf.pages, layers["fields"]):
                field = fields[field_id]
                style = report.styles.get(field_id)
                styled = (
                    field.model_copy(update=style.model_dump(exclude_none=True))
                    if style
                    else field
                )
                box = report.layout.get(field_id, field)
                page = definition.pages[box.page]
                for line in layer_page.extract_text_lines(return_chars=True):
                    identity += 1
                    size = line["chars"][0]["size"]
                    x = box.x * page.width + line["x0"]
                    y = box.y * page.height + line["top"] - size * 0.12
                    _text_box(
                        page_paragraphs[box.page],
                        line,
                        x,
                        y,
                        max(1, page.width - x),
                        identity,
                        styled,
                    )
    output = BytesIO()
    document.save(output)
    return output.getvalue()
