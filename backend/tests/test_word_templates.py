"""Word conversion, retention, fixed-layout exports and revision isolation."""

from io import BytesIO
import json
import shutil
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree
from pypdf import PdfReader
import pytest

from app import template_main
from app.template_models import TemplateReport
from app.template_pdf import clean_pdf
from app.template_word import generate_docx
from app.word_conversion import (
    validate_docx,
    docx_to_pdf,
    DOCX_MIME,
    ConverterUnavailable,
)
from test_templates import definition, sample_pdf
from test_audit import create_shared


def sample_word():
    document = Document()
    document.add_heading("Inspection report", level=0)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, label in zip(table.rows, ["Company", "Location", "Remarks"]):
        row.cells[0].text = label
        row.cells[1].text = "Sample text"
    document.add_page_break()
    document.add_paragraph("Second page reference")
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def altered_word(name, value):
    output = BytesIO()
    with (
        ZipFile(BytesIO(sample_word())) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as archive,
    ):
        for item in original.infolist():
            if item.filename != name:
                archive.writestr(item.filename, original.read(item))
        archive.writestr(name, value)
    return output.getvalue()


def docx_text(payload):
    with ZipFile(BytesIO(payload)) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        return root, " ".join(root.xpath("//*[local-name()='t']/text()"))


def test_valid_word_and_invalid_archives():
    validate_docx(sample_word())
    for payload in (b"not a Word document", b"PKfake", sample_pdf()):
        with pytest.raises(ValueError, match="Word"):
            validate_docx(payload)


@pytest.mark.parametrize(
    "name,value",
    [
        ("word/vbaProject.bin", b"macro"),
        ("word/embeddings/file.bin", b"embedded"),
        ("word/activeX/file.xml", b"<root/>"),
        (
            "word/_rels/header1.xml.rels",
            b'<Relationships><Relationship TargetMode="External" Type="image" Target="http://169.254.169.254/metadata"/></Relationships>',
        ),
        (
            "word/document.xml",
            b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x/>',
        ),
        (
            "word/document.xml",
            b'<root><instrText>INCLUDETEXT "https://example.com"</instrText></root>',
        ),
        (
            "word/header1.xml",
            b'<root><instrText>INCLUDEPICTURE "https://example.com"</instrText></root>',
        ),
        ("word/document.xml", b"<root><altChunk/></root>"),
        ("word/document.xml", b"<unclosed>"),
    ],
)
def test_rejects_unsafe_word_features(name, value):
    with pytest.raises(ValueError):
        validate_docx(altered_word(name, value))


def test_zip_expansion_limit(monkeypatch):
    from app import word_conversion

    monkeypatch.setattr(word_conversion, "MAX_EXPANDED", 100)
    with pytest.raises(ValueError, match="processing limit"):
        validate_docx(sample_word())


def test_missing_converter_explained(monkeypatch):
    monkeypatch.delenv("LIBREOFFICE_BINARY", raising=False)
    monkeypatch.setattr(shutil, "which", lambda _: None)
    with pytest.raises(ConverterUnavailable, match="Rebuild"):
        docx_to_pdf(sample_word())


def test_word_analyze_save_revision_and_original_retention(api, monkeypatch):
    original = sample_word()
    monkeypatch.setattr(template_main, "docx_to_pdf", lambda payload: sample_pdf())
    files = {"pdf": ("job.docx", original, DOCX_MIME)}
    analysis = api.post("/api/templates/analyze", files=files)
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()["source_format"] == "docx"
    assert analysis.json()["render_pdf"]
    saved = api.post(
        "/api/templates",
        files=files,
        data={"title": "Word job", "definition_json": definition().model_dump_json()},
    )
    assert saved.status_code == 200, saved.text
    detail = saved.json()
    base = "/api/templates/" + detail["id"]
    assert detail["source_format"] == "docx"
    assert api.get(base + "/source").content.startswith(b"%PDF-")
    assert api.get(base + "/original").content == original
    draft = api.post(
        base + "/drafts",
        data={"report_json": '{"title":"Retained Word", "values":{"company":"Acme"}}'},
    ).json()
    revision = api.put(
        base,
        data={
            "title": "Renamed",
            "definition_json": definition().model_dump_json(),
            "expected_revision_id": detail["revision_id"],
        },
    )
    assert revision.status_code == 200, revision.text
    assert revision.json()["source_format"] == "docx"
    assert api.get(base + "/original").content == original
    for output in ("pdf", "docx", "original"):
        result = api.get(base + f"/drafts/{draft['id']}/download?format={output}")
        assert result.status_code == 200, result.text[:500]
        assert result.content.startswith(b"%PDF-" if output == "pdf" else b"PK")
    # Switching source format cannot change the original-format export of old drafts.
    replaced = api.put(
        base,
        files={"pdf": ("replacement.pdf", sample_pdf(), "application/pdf")},
        data={
            "title": "PDF job",
            "definition_json": definition().model_dump_json(),
            "expected_revision_id": revision.json()["revision_id"],
        },
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["source_format"] == "pdf"
    assert api.get(
        base + f"/drafts/{draft['id']}/download?format=original"
    ).content.startswith(b"PK")
    assert api.post(
        base + "/generate?format=original", data={"report_json": '{"title":"New PDF"}'}
    ).content.startswith(b"%PDF-")
    assert api.delete(base).status_code == 204
    assert api.delete(base + "/permanent?confirmed=true").status_code == 204
    assert api.get(base + "/original").status_code == 404


def test_pdf_to_word_has_native_text_styling_and_page_geometry(api):
    base = create_shared(api)
    form = {
        "title": "Word export",
        "values": {"company": "Acme\nBranch", "temp": "25"},
        "styles": {
            "company": {
                "font": "Helvetica-Bold",
                "color": "#008800",
                "alignment": "right",
            }
        },
    }
    response = api.post(
        base + "/generate?format=docx", data={"report_json": json.dumps(form)}
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == DOCX_MIME
    assert 'filename="Word export.docx"' in response.headers["content-disposition"]
    root, text = docx_text(response.content)
    assert "Acme" in text and "Branch" in text and "25" in text
    assert len(root.xpath("//*[local-name()='sectPr']")) == 2
    assert (
        root.xpath("//*[local-name()='color']/@*[local-name()='val']").count("008800")
        == 2
    )
    assert len(root.xpath("//*[local-name()='anchor']")) == 2
    assert len(Document(BytesIO(response.content)).sections) == 2
    assert (
        api.post(base + "/generate?format=exe", data={"report_json": "{}"}).status_code
        == 422
    )
    assert (
        api.post(
            base + "/generate?format=docx",
            data={"report_json": '{"values":{"unknown":"no"}}'},
        ).status_code
        == 422
    )


def test_word_export_handles_images_shapes_calculations_and_blanks():
    from PIL import Image

    schema = definition()
    blank, _ = clean_pdf(sample_pdf(), schema)
    image = BytesIO()
    Image.new("RGB", (100, 60), "orange").save(image, format="PNG")
    report = TemplateReport(
        values={"company": "Native value", "temp": "25"},
        rectangles=[
            dict(id="tick", page=0, kind="tick", x=0.7, y=0.05, width=0.05, height=0.05)
        ],
    )
    result = generate_docx(blank, schema, report, {"photo": image.getvalue()})
    root, text = docx_text(result)
    assert "Native value" in text
    with ZipFile(BytesIO(result)) as archive:
        assert (
            len([name for name in archive.namelist() if name.startswith("word/media/")])
            == 2
        )
    assert generate_docx(blank, schema, TemplateReport(), {}).startswith(b"PK")


@pytest.mark.skipif(
    not shutil.which("soffice"),
    reason="LibreOffice integration runs inside the service image",
)
def test_real_word_conversion_and_export_roundtrip():
    converted = docx_to_pdf(sample_word())
    assert len(PdfReader(BytesIO(converted)).pages) == 2
    schema = definition()
    blank, _ = clean_pdf(sample_pdf(), schema)
    exported = generate_docx(
        blank,
        schema,
        TemplateReport(values={"company": "Native Word value", "temp": "42"}),
        {},
    )
    roundtrip = PdfReader(BytesIO(docx_to_pdf(exported)))
    assert len(roundtrip.pages) == 2
    assert "Native Word value" in " ".join(
        page.extract_text() for page in roundtrip.pages
    )
